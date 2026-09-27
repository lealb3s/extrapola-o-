"""
Exporta o caso atual para uma planilha Excel COM FÓRMULAS (não só valores):
alterar qualquer célula azul recalcula toda a extrapolação e os gráficos.

Abas:
  - "Extrapolação": parâmetros, tabela Froude × Hughes por velocidade e gráficos
    de R_T e P_E versus Fn.
  - "Modelo": nome, referência, características e imagem do modelo.
"""

from __future__ import annotations

import io
from typing import Optional

from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .extrapolacao import Parametros, NO_EM_MS, HP_EM_W
from .atrito import coef_atrito

AZUL = Font(name="Arial", color="0000FF")
PRETO = Font(name="Arial")
NEGRITO = Font(name="Arial", bold=True)
TITULO = Font(name="Arial", bold=True, size=14)
CAB = Font(name="Arial", bold=True, color="FFFFFF")
FUNDO_CAB = PatternFill("solid", fgColor="1F4E78")
FUNDO_FROUDE = PatternFill("solid", fgColor="DDEBF7")
FUNDO_HUGHES = PatternFill("solid", fgColor="FCE4D6")
FUNDO_ENTRADA = PatternFill("solid", fgColor="FFF2CC")
FINA = Side(style="thin", color="BFBFBF")
BORDA = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)


def _cf_formula(linha: str, re_ref: str) -> Optional[str]:
    if linha == "ITTC-1957":
        return f"=0.075/(LOG10({re_ref})-2)^2"
    if linha == "Hughes (1954)":
        return f"=0.066/(LOG10({re_ref})-2.03)^2"
    return None  # Schoenherr é implícita: exporta valor


def gerar_xlsx(
    p: Parametros,
    V_m,
    R_Tm,
    nome: str = "",
    referencia: str = "",
    caracteristicas: str = "",
    imagem: Optional[bytes] = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Extrapolação"

    ws["A1"] = "Extrapolação modelo-protótipo — Métodos de Froude e Hughes"
    ws["A1"].font = TITULO
    ws["A2"] = (nome or "Caso sem nome") + ("  |  Ref.: " + referencia if referencia else "")
    ws["A2"].font = Font(name="Arial", italic=True)
    ws["A3"] = "Legenda: células AZUIS com fundo amarelo são entradas (edite-as); o resto são fórmulas."
    ws["A3"].font = Font(name="Arial", italic=True, color="7F7F7F")

    # ---- Bloco de parâmetros --------------------------------------------------
    linhas = [
        ("MODELO", None, None, None),
        ("L_m", p.L_m, "m", "Comprimento do modelo"),
        ("S_m", p.S_m, "m²", "Superfície molhada do modelo"),
        ("ρ_m", p.rho_m, "kg/m³", "Massa específica — água do tanque"),
        ("ν_m", p.nu_m, "m²/s", "Viscosidade cinemática — água do tanque"),
        ("PROTÓTIPO", None, None, None),
        ("L_s", p.L_s, "m", "Comprimento do navio"),
        ("S_s", None, "m²", "Superfície molhada do navio"),
        ("ρ_s", p.rho_s, "kg/m³", "Massa específica — água do mar"),
        ("ν_s", p.nu_s, "m²/s", "Viscosidade cinemática — água do mar"),
        ("MÉTODO", None, None, None),
        ("k", p.k, "-", "Fator de forma (Hughes)"),
        ("g", p.g, "m/s²", "Aceleração da gravidade"),
        ("C_A", p.C_A, "-", "Correção de correlação (0 no exercício)"),
        ("λ", None, "-", "Escala = L_s / L_m"),
    ]
    ref = {}
    r0 = 4
    for i, (rot, val, un, desc) in enumerate(linhas):
        r = r0 + i
        c = ws.cell(r, 1, rot)
        if val is None and un is None:
            c.font = NEGRITO
            continue
        c.font = PRETO
        ws.cell(r, 3, un).font = PRETO
        ws.cell(r, 4, desc).font = Font(name="Arial", color="7F7F7F")
        ref[rot] = f"$B${r}"
        cel = ws.cell(r, 2)
        if val is not None:
            cel.value = val
            cel.font = AZUL
            cel.fill = FUNDO_ENTRADA
        cel.border = BORDA
    # S_s e λ
    ws[ref["λ"].replace("$", "")] = f"={ref['L_s']}/{ref['L_m']}"
    ws[ref["λ"].replace("$", "")].font = PRETO
    ws[ref["λ"].replace("$", "")].number_format = "0.0000"
    cs = ws[ref["S_s"].replace("$", "")]
    if p.S_s is None:
        cs.value = f"={ref['S_m']}*{ref['λ']}^2"
        cs.font = PRETO
        cs.comment = Comment("Calculada por semelhança: S_s = S_m·λ². Pode ser substituída por um valor.", "app")
    else:
        cs.value = p.S_s
        cs.font = AZUL
        cs.fill = FUNDO_ENTRADA
        cs.comment = Comment("Valor informado (fonte do caso). Por semelhança seria S_m·λ².", "app")
    for rot in ("ν_m", "ν_s"):
        ws[ref[rot].replace("$", "")].number_format = "0.000E+00"
    ws.cell(r0 + len(linhas), 1, f"Linha de atrito: {p.linha_atrito}").font = Font(name="Arial", italic=True)

    # ---- Tabela por velocidade -----------------------------------------------
    cab_r = r0 + len(linhas) + 2          # linha do grupo
    hdr = cab_r + 1                       # linha dos cabeçalhos
    d0 = hdr + 1                          # primeira linha de dados
    colunas = [
        ("V_m [m/s]", "0.000"), ("R_Tm [N]", "0.000"), ("Fn", "0.0000"), ("Re_m", "0.000E+00"),
        ("C_Tm", "0.000E+00"), ("C_Fm", "0.000E+00"), ("V_s [m/s]", "0.000"), ("V_s [nós]", "0.00"),
        ("Re_s", "0.000E+00"), ("C_Fs", "0.000E+00"),
        ("C_R", "0.000E+00"), ("C_Ts", "0.000E+00"), ("R_Ts [kN]", "#,##0.00"),
        ("P_E [kW]", "#,##0.0"), ("P_E [hp]", "#,##0.0"),
        ("C_Vm", "0.000E+00"), ("C_W", "0.000E+00"), ("C_Vs", "0.000E+00"), ("C_Ts", "0.000E+00"),
        ("R_Ts [kN]", "#,##0.00"), ("P_E [kW]", "#,##0.0"), ("P_E [hp]", "#,##0.0"),
    ]
    ws.cell(cab_r, 1, "Ensaio do modelo e semelhança").font = NEGRITO
    ws.cell(cab_r, 11, "(A) Froude — sem fator de forma").font = NEGRITO
    ws.cell(cab_r, 16, "(B) Hughes — com fator de forma").font = NEGRITO
    for j, (h, _) in enumerate(colunas, start=1):
        c = ws.cell(hdr, j, h)
        c.font = CAB
        c.fill = FUNDO_CAB
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = BORDA

    L_m, S_m, rho_m, nu_m = ref["L_m"], ref["S_m"], ref["ρ_m"], ref["ν_m"]
    L_s, S_s, rho_s, nu_s = ref["L_s"], ref["S_s"], ref["ρ_s"], ref["ν_s"]
    k, g, CA, lam = ref["k"], ref["g"], ref["C_A"], ref["λ"]

    n = len(V_m)
    for i in range(n):
        r = d0 + i
        f = {}
        f["A"] = float(V_m[i])
        f["B"] = float(R_Tm[i])
        f["C"] = f"=A{r}/SQRT({g}*{L_m})"
        f["D"] = f"=A{r}*{L_m}/{nu_m}"
        f["E"] = f"=B{r}/(0.5*{rho_m}*{S_m}*A{r}^2)"
        f["F"] = _cf_formula(p.linha_atrito, f"D{r}")
        f["G"] = f"=A{r}*SQRT({lam})"
        f["H"] = f"=G{r}/{NO_EM_MS}"
        f["I"] = f"=G{r}*{L_s}/{nu_s}"
        f["J"] = _cf_formula(p.linha_atrito, f"I{r}")
        f["K"] = f"=E{r}-F{r}"
        f["L"] = f"=J{r}+K{r}+{CA}"
        f["M"] = f"=L{r}*0.5*{rho_s}*{S_s}*G{r}^2/1000"
        f["N"] = f"=M{r}*G{r}"
        f["O"] = f"=N{r}*1000/{HP_EM_W}"
        f["P"] = f"=(1+{k})*F{r}"
        f["Q"] = f"=E{r}-P{r}"
        f["R"] = f"=(1+{k})*J{r}"
        f["S"] = f"=R{r}+Q{r}+{CA}"
        f["T"] = f"=S{r}*0.5*{rho_s}*{S_s}*G{r}^2/1000"
        f["U"] = f"=T{r}*G{r}"
        f["V"] = f"=U{r}*1000/{HP_EM_W}"
        if f["F"] is None:  # Schoenherr → valores
            f["F"] = float(coef_atrito(float(V_m[i]) * p.L_m / p.nu_m, p.linha_atrito))
            f["J"] = float(coef_atrito(float(V_m[i]) * (p.escala ** 0.5) * p.L_s / p.nu_s, p.linha_atrito))
        for j, (_, fmt) in enumerate(colunas, start=1):
            col = get_column_letter(j)
            c = ws.cell(r, j, f[col])
            c.number_format = fmt
            c.border = BORDA
            c.font = AZUL if col in ("A", "B") else PRETO
            if col in ("A", "B"):
                c.fill = FUNDO_ENTRADA
            elif 11 <= j <= 15:
                c.fill = FUNDO_FROUDE
            elif j >= 16:
                c.fill = FUNDO_HUGHES
    if p.linha_atrito not in ("ITTC-1957", "Hughes (1954)"):
        ws.cell(d0 + n, 1, "Obs.: C_F de Schoenherr é implícita — colunas C_Fm e C_Fs exportadas como valores.").font = Font(name="Arial", italic=True, color="C00000")

    for j in range(1, len(colunas) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 12.5
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.row_dimensions[hdr].height = 30
    ws.freeze_panes = ws.cell(d0, 3)

    # ---- Gráficos ----------------------------------------------------------------
    if n >= 1:
        r1, r2 = d0, d0 + n - 1
        xref = Reference(ws, min_col=3, min_row=r1, max_row=r2)

        def grafico(titulo, ytit, col_f, col_h):
            ch = ScatterChart()
            ch.title = titulo
            ch.style = 13
            ch.x_axis.title = "Número de Froude, Fn"
            ch.y_axis.title = ytit
            ch.height, ch.width = 8, 15
            for col, nome_s, cor in ((col_f, "Froude", "2E6FBA"), (col_h, "Hughes", "D9822B")):
                s = Series(Reference(ws, min_col=col, min_row=r1, max_row=r2), xref, title=nome_s)
                s.marker.symbol = "circle"
                s.marker.graphicalProperties.solidFill = cor
                s.marker.graphicalProperties.line.solidFill = cor
                s.graphicalProperties.line.solidFill = cor
                s.graphicalProperties.line.width = 28575
                s.smooth = False
                ch.series.append(s)
            ch.x_axis.delete = False
            ch.y_axis.delete = False
            return ch

        pos = d0 + n + 2
        ws.add_chart(grafico("Resistência total do navio R_Ts × Fn", "R_Ts [kN]", 13, 20), f"A{pos}")
        ws.add_chart(grafico("Potência efetiva P_E × Fn", "P_E [kW]", 14, 21), f"I{pos}")

    # ---- Aba do modelo -------------------------------------------------------------
    wm = wb.create_sheet("Modelo")
    wm["A1"] = "Modelo selecionado"
    wm["A1"].font = TITULO
    info = [
        ("Nome", nome),
        ("Referência", referencia),
        ("Características", caracteristicas),
        ("Escala λ", f"='{ws.title}'!{lam}"),
        ("L_m [m]", f"='{ws.title}'!{L_m}"),
        ("S_m [m²]", f"='{ws.title}'!{S_m}"),
        ("L_s [m]", f"='{ws.title}'!{L_s}"),
        ("S_s [m²]", f"='{ws.title}'!{S_s}"),
        ("k", f"='{ws.title}'!{k}"),
        ("Linha de atrito", p.linha_atrito),
    ]
    for i, (a, b) in enumerate(info, start=3):
        wm.cell(i, 1, a).font = NEGRITO
        c = wm.cell(i, 2, b)
        c.font = PRETO
        c.alignment = Alignment(wrap_text=True, vertical="top")
    wm.column_dimensions["A"].width = 18
    wm.column_dimensions["B"].width = 70
    if imagem:
        try:
            from openpyxl.drawing.image import Image as XLImage
            from PIL import Image as PILImage

            pil = PILImage.open(io.BytesIO(imagem))
            pil.thumbnail((640, 420))
            buf = io.BytesIO()
            pil.convert("RGB").save(buf, format="PNG")
            buf.seek(0)
            wm.add_image(XLImage(buf), "A15")
        except Exception:  # imagem inválida não impede a exportação
            wm["A15"] = "(não foi possível incorporar a imagem enviada)"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
